from __future__ import annotations
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from app.digest.common import AUDIENCE, detail_line, fmt_date, group_by_topic, key_figures
from app.models import DigestEvent, UserProfile


def _hyperlink(par, text: str, url: str):
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rpr.append(c)
    r.append(rpr); t = OxmlElement("w:t"); t.text = text; r.append(t); h.append(r); par._p.append(h)


def export_docx(events: list[DigestEvent], profile: UserProfile, path: str | Path, meta: dict | None = None) -> Path:
    meta = meta or {}
    doc = Document(); sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.65); sec.left_margin = sec.right_margin = Inches(.75)
    doc.styles["Normal"].font.name = "Arial"; doc.styles["Normal"].font.size = Pt(10)
    kind = {"Месяц": "Ежемесячный", "День": "Ежедневный"}.get(meta.get("period_label"), "")
    h = doc.add_heading(f"{kind} дайджест ИТ/ИБ".strip(), 0); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{AUDIENCE.get(profile.role, profile.role)} · {(meta.get('generated_at') or datetime.now()).strftime('%d.%m.%Y %H:%M')}"
              + (f" · период: {meta['period_label']}" if meta.get("period_label") else ""))
    doc.add_paragraph(f"Интересы: {', '.join(profile.interests) or 'не заданы'}")
    if profile.role == "cio" or meta.get("period_label") == "Месяц":
        k = key_figures(events)
        doc.add_heading("Ключевые цифры", 1)
        doc.add_paragraph(f"Событий: {k['events']} (публикаций: {k['publications']}); уязвимости: {k['vulnerabilities']}; "
                          f"атаки: {k['attacks']}; с exploit/PoC: {k['with_exploit']}")
    if not events:
        doc.add_paragraph("По заданным интересам и периоду релевантных материалов не найдено.")
    for topic, items in group_by_topic(events).items():
        doc.add_heading(topic, 1)
        for e in items:
            doc.add_heading(e.title, 2)
            m = doc.add_paragraph(); r = m.add_run(f"{e.category} · {fmt_date(e.published_at)} · релевантность {e.score:.0%}")
            r.italic = True; r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            doc.add_paragraph(e.summary)
            w = doc.add_paragraph(); w.add_run("Почему важно: ").bold = True; w.add_run(e.why_it_matters)
            if detail_line(e):
                d = doc.add_paragraph(); d.add_run("Детали: ").bold = True; d.add_run(detail_line(e))
            if e.evidence:
                q = doc.add_paragraph(); q.paragraph_format.left_indent = Inches(.3)
                q.add_run("«" + e.evidence + "»").italic = True
            doc.add_paragraph().add_run("Источники:").bold = True
            for l in e.links or [{"source": "", "title": s, "url": s} for s in e.sources]:
                b = doc.add_paragraph(style="List Bullet")
                _hyperlink(b, f"{l.get('source', '')}: {l['title']}".strip(": "), l["url"])
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True); doc.save(path); return path
